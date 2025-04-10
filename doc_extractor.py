import os
import json
import fitz
import pdfplumber 
import pandas as pd
import docx
import tiktoken
from typing import List, Dict

def count_tokens(text: str, tokenizer) -> int:
    """Count tokens using cl100k_base tokenizer."""
    return len(tokenizer.encode(text))

def split_text(text, max_tokens, tokenizer, base_filename, base_chunk_id):
    """Splits text into smaller chunks if it exceeds max_tokens."""
    chunks = []
    encoded_tokens = tokenizer.encode(text)

    for i in range(0, len(encoded_tokens), max_tokens):
        chunk_tokens = encoded_tokens[i:i+max_tokens]
        chunk_text = tokenizer.decode(chunk_tokens)
        chunks.append({
            "file": base_filename,
            "type": "text",
            "text": chunk_text,
            "chunk_id": f"{base_chunk_id}_part{i//max_tokens}",
            "tokens": len(chunk_tokens)
        })
    
    return chunks

def extract_text_from_docx(file_path: str, tokenizer, max_tokens=200):
    """Extract and split text from DOCX files."""
    doc = docx.Document(file_path)
    extracted = []
    base_filename = os.path.basename(file_path)

    chunk_id = 0
    for para in doc.paragraphs:
        if para.text.strip():
            extracted.extend(split_text(para.text.strip(), max_tokens, tokenizer, base_filename, f"{base_filename}_chunk{chunk_id}"))
            chunk_id += 1

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                extracted.extend(split_text(row_text, max_tokens, tokenizer, base_filename, f"{base_filename}_chunk{chunk_id}"))
                chunk_id += 1

    return extracted

def extract_text_from_pdf(file_path: str, tokenizer, max_tokens=200):
    """Extract and split text from PDFs."""
    extracted = []
    base_filename = os.path.basename(file_path)

    with fitz.open(file_path) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            if text:
                extracted.extend(split_text(text, max_tokens, tokenizer, base_filename, f"{base_filename}_page{page_num}"))

    with pdfplumber.open(file_path) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            tables = page.extract_tables()
            chunk_id = 1
            for table in tables:
                for row in table:
                    row_text = " | ".join(cell if cell else "" for cell in row)
                    if row_text.strip():
                        extracted.extend(split_text(row_text, max_tokens, tokenizer, base_filename, f"{base_filename}_page{page_num}_table{chunk_id}"))
                        chunk_id += 1

    return extracted

def extract_text_from_csv(file_path: str, tokenizer) -> List[Dict]:
    """Extract text from a CSV file."""
    df = pd.read_csv(file_path, dtype=str, encoding="utf-8", keep_default_na=False)
    extracted = []
    base_filename = os.path.basename(file_path)

    for i, row in df.iterrows():
        row_text = " | ".join(row.dropna().astype(str).values)
        if row_text.strip():
            extracted.append({
                "file": base_filename,
                "type": "table",
                "text": row_text,
                "chunk_id": f"{base_filename}_row{i}",
                "tokens": count_tokens(row_text, tokenizer)
            })

    return extracted

def extract_text_from_excel(file_path: str, tokenizer) -> List[Dict]:
    """Extract text from an Excel file."""
    xls = pd.ExcelFile(file_path)
    extracted = []
    base_filename = os.path.basename(file_path)

    for sheet_name in xls.sheet_names:
        df = pd.read_excel(xls, sheet_name=sheet_name, dtype=str, keep_default_na=False)
        for i, row in df.iterrows():
            row_text = " | ".join(row.dropna().astype(str).values)
            if row_text.strip():
                extracted.append({
                    "file": base_filename,
                    "sheet": sheet_name,
                    "type": "table",
                    "text": row_text,
                    "chunk_id": f"{base_filename}_{sheet_name}_row{i}",
                    "tokens": count_tokens(row_text, tokenizer)
                })

    return extracted

def extract_text(file_path: str, tokenizer) -> List[Dict]:
    """Extract text from different document types."""
    ext = os.path.splitext(file_path)[-1].lower()
    
    if ext == ".docx":
        return extract_text_from_docx(file_path, tokenizer)
    elif ext == ".pdf":
        return extract_text_from_pdf(file_path, tokenizer)
    elif ext == ".csv":
        return extract_text_from_csv(file_path, tokenizer)
    elif ext in [".xlsx", ".xls", ".xlsm"]:
        return extract_text_from_excel(file_path, tokenizer)
    else:
        raise ValueError(f"Unsupported file format: {ext}")

def process_folder(folder_path: str, output_json: str):
    """Process all files in a folder and save extracted text as JSON."""
    tokenizer = tiktoken.get_encoding("cl100k_base")
    all_extracted = []

    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if os.path.isfile(file_path):
            try:
                extracted_text = extract_text(file_path, tokenizer)
                all_extracted.extend(extracted_text)
            except Exception as e:
                print(f" ERROR PROCESSING {filename}: {e} !!!!!!!!!!!!!!")

    with open(output_json, "w", encoding="utf-8") as f:
        json.dump(all_extracted, f, indent=4, ensure_ascii=False)

    print(f"----------------------- Extraction complete! Data saved to: {output_json} -----------------------")

if __name__ == "__main__":
    folder_path = "Dr.X Files/"
    output_json = "extracted_text.json"
    process_folder(folder_path, output_json)
